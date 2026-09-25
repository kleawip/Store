import build_feature_expansion_addendum as b
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def section(doc, title, intro=None):
    b.add_heading(doc, title, level=2)
    if intro:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(intro)
        b.set_repeatable_font(r, size=8.8, color=b.MUTED)


def build():
    b.OUT.parent.mkdir(parents=True, exist_ok=True)
    b.LOGO_PNG.parent.mkdir(parents=True, exist_ok=True)
    doc = b.configure_document()
    b.add_cover(doc)
    b.add_heading(doc, "Expansion areas", level=2)
    b.add_table(
        doc,
        ["Service", "After sales", "Delivery", "Discovery", "Retention and control"],
        [["Tickets and SLA", "Returns, exchanges and refunds", "NDR and RTO", "Reviews, search and GSM finder", "Store credit, referrals and exceptions"]],
        widths=[5.3, 5.3, 5.3, 5.3, 5.3],
        font_size=7.8,
        header_fill=b.PURPLE,
    )

    b.new_page(
        doc,
        "01  Purpose and scale",
        "New capabilities beyond the original ecommerce baseline",
        "This addendum expands selected operational and customer-experience areas identified after the original scope was shared. It does not replace, reduce or re-price the original proposal, and it does not repeat baseline functionality as a new deliverable.",
    )
    b.add_metrics(
        doc,
        [("67 to 82", "Customer screens"), ("44+ to 73+", "Admin views"), ("111+ to 155+", "Combined views"), ("+44", "New destinations")],
        fill=b.PALE_PURPLE,
    )
    b.add_table(
        doc,
        ["Area", "Original baseline", "Expansion in this addendum"],
        [
            ["Bulk purchasing", "Bulk / Wholesale Enquiry page is already in baseline", "Confirm quote-only or online bulk order, MOQ and tiered terms; no duplicate page count"],
            ["Support", "Ticket destination and support role", "Service desk, ownership, SLA, conversations, escalation and analytics"],
            ["Client showcase", "No dedicated destination in original page list", "One new Clients We Work With page using approved client names, logos and references"],
            ["After sales", "Return, pickup, inspection and refund/replacement", "Policy rules, exchanges, COD refunds, store credit and approval levels"],
            ["Delivery", "Tracking, NDR status and RTO reporting", "NDR action queue, customer response, RTO risk rules and cost control"],
            ["Trust and discovery", "Reviews, search and filters", "Moderation, Q&A, merchandising, comparison and guided GSM selection"],
            ["Retention", "Reorder, cohorts and reminders", "Back-in-stock, store credit, referral and stronger communication controls"],
            ["Governance", "Roles, logs, backups and integrations", "Exception queues, high-risk approvals, privacy requests and sessions"],
        ],
        widths=[4.1, 9.2, 13.3],
        font_size=7.7,
    )
    b.add_body(doc, "Counts represent major destinations and workflows. Tabs, modals, empty states, errors, permissions and light/dark responsive states are not separately counted.", bold_lead="Counting method. ")

    b.new_page(
        doc,
        "02  Service and after sales",
        "Support, returns and refunds become connected operations",
        "The customer sees a clear case history. Kleawip staff receive ownership, policy control, approval steps and measurable resolution workflows connected to the relevant order, payment, shipment and item.",
    )
    b.add_table(
        doc,
        ["Customer support and service desk", "Advanced returns, exchanges and refunds"],
        [[
            "Support centre and guided issue selection\nMy Tickets list and ticket conversation\nOrder, payment, shipment, return or product linkage\nValidated attachments and customer-visible status\nAgent/team assignment and ownership history\nPriority, first-response and resolution SLA\nInternal notes, reply templates and escalation\nWorkload, ageing, reopen and satisfaction reports",
            "Eligibility by product, delivery date, reason and condition\nItem-level return, replacement, exchange or partial refund\nRMA reference, evidence and structured reason codes\nReverse pickup, AWB and pickup exception handling\nWarehouse QC and stock disposition\nOnline refund status and reconciliation\nControlled COD refund details and approval\nStore credit or coupon compensation where approved\nReturn, refund and product-quality reporting",
        ]],
        widths=[13.3, 13.3],
        font_size=8.25,
    )
    section(doc, "AI assisted support with human approval")
    b.add_bullets(doc, [
        "Summarise long ticket threads and highlight the unresolved issue.",
        "Suggest category, priority and an editable response using approved Kleawip knowledge.",
        "Keep private order and customer data permission-based and logged.",
        "Require authorised staff approval for every refund, cancellation, replacement and account change.",
    ], columns=2, size=8.5)
    b.add_body(doc, "Refund initiation, provider acknowledgement, final processing and reconciliation remain separate states so staff never mark a refund complete from an uncertain response.", bold_lead="Financial control. ")

    b.new_page(
        doc,
        "03  Delivery and trust",
        "NDR, RTO, verified reviews and product questions",
        "Courier failures and customer feedback become actionable records instead of isolated provider-dashboard entries, chat messages or unverified website content.",
    )
    section(doc, "NDR and RTO control centre")
    b.add_table(
        doc,
        ["Workflow", "Expanded control", "Outcome"],
        [
            ["NDR intake", "Reason, courier status, owner, response deadline and next action", "One queue across shipments"],
            ["Customer response", "Secure confirmation of address, date, reattempt or cancellation where supported", "Faster recovery"],
            ["RTO risk", "Rules using pincode, order value, history and prior RTO", "Prepaid or partial-COD control"],
            ["Reconciliation", "Forward/return freight, advance, COD due and remittance", "True failed-order cost"],
        ],
        widths=[4.0, 14.3, 8.3],
        font_size=7.8,
    )
    section(doc, "Verified reviews and product questions")
    b.add_table(
        doc,
        ["Capability", "Customer experience", "Administration control"],
        [
            ["Reviews", "Rating, text and optional photo/video after delivery", "Verified-purchase link, moderation and history"],
            ["Merchant reply", "Visible Kleawip response", "Authorised reply and edit history"],
            ["Product Q&A", "Ask and browse product/use questions", "Assignment, answer, approval and publishing"],
            ["Trust protection", "Report inappropriate content", "Spam signals, rights checks and auditable decision"],
        ],
        widths=[4.0, 11.2, 11.4],
        font_size=7.8,
    )
    b.add_body(doc, "Only genuine approved feedback contributes to displayed ratings. Automated courier actions remain dependent on the selected provider's API and account capabilities.", bold_lead="Control boundary. ")

    b.new_page(
        doc,
        "04  Discovery and retention",
        "Help customers choose correctly and return for useful reasons",
        "The expansion improves how shoppers find a suitable product and how Kleawip follows up after genuine interest or purchase, subject to approved consent and communication rules.",
    )
    section(doc, "Advanced search and guided selection")
    b.add_table(
        doc,
        ["Capability", "Experience", "Client control"],
        [
            ["Search language", "Typo tolerance, normalization and synonyms", "Approved terms and priorities"],
            ["Merchandising", "Relevant promoted items without hiding stronger matches", "Rules by query, category, stock and period"],
            ["Zero results", "Alternative terms, categories and products", "Zero-result query report"],
            ["Comparison", "Side-by-side GSM, size, weave, use, absorbency and care", "Comparable fields and eligible products"],
            ["GSM finder", "Guided questions about surface, drying need and size", "Questions, rules and recommendations"],
        ],
        widths=[4.0, 12.4, 10.2],
        font_size=7.7,
    )
    section(doc, "Bulk purchase and wholesale enquiry")
    b.add_body(doc, "The original 67-page scope already includes a Bulk / Wholesale Enquiry page. The enhancement defines its business workflow: product and variant interest, estimated quantity, business contact, delivery location, requested timeline and enquiry follow-up. Whether Kleawip needs quote-only enquiries or a separate bulk checkout with MOQ and tiered pricing must be confirmed in Phase 0; this page is not counted twice.")
    section(doc, "Retention and account value")
    b.add_table(
        doc,
        ["Feature", "Customer value", "Safeguard"],
        [
            ["Back in stock", "Subscribe to an unavailable variant", "Frequency cap and unsubscribe"],
            ["Abandoned cart", "Secure return to the saved cart", "Delay, exclusions and consent"],
            ["Store credit", "Use an approved account balance", "Immutable ledger, expiry and permissions"],
            ["Referral centre", "Share referral and view rewards", "Qualification, fraud and reversal rules"],
            ["Segments", "More relevant offers and messages", "Consent, exclusion and export controls"],
        ],
        widths=[4.0, 12.4, 10.2],
        font_size=7.7,
    )

    b.new_page(
        doc,
        "05  Communication and reliability",
        "Every failed or uncertain operation receives a safe next action",
        "The baseline notification system is extended with administrative control and central exception handling for payments, refunds, shipments, inventory and messages.",
    )
    section(doc, "Notification operations")
    b.add_bullets(doc, [
        "Editable approved templates with safe variables and preview data.",
        "Version history, publishing approval and rollback.",
        "Delivery logs by order, customer, channel, event and provider status.",
        "Automatic retries with duplicate-send protection.",
        "Authorised replay with reason and audit history.",
        "Preference-centre and channel unsubscribe handling where applicable.",
    ], columns=2, size=8.4)
    section(doc, "Operations exception queue")
    b.add_table(
        doc,
        ["Area", "Failures made visible", "Resolution record"],
        [
            ["Payment", "Late success, duplicate callback, signature or amount mismatch", "Verification and order decision"],
            ["Refund", "Pending, failed, processed or reconciliation mismatch", "Reference, provider status and approval"],
            ["Shipping", "AWB, pickup, tracking or cancellation failure", "Response, retry and manual action"],
            ["Inventory", "Reservation expiry or adjustment conflict", "Ledger movement and resulting balance"],
            ["Notification", "Rejected, bounced, delayed or unavailable", "Template, status, retry and replay"],
        ],
        widths=[3.6, 13.8, 9.2],
        font_size=7.8,
    )
    b.add_body(doc, "Manual refunds, store-credit issues, stock write-offs and bulk rewards use maker-checker approval and a permanent audit entry.", bold_lead="High-risk actions. ")

    b.new_page(
        doc,
        "06  Privacy and scope separation",
        "Stronger customer-data workflows without mixing service categories",
        "The expansion adds practical controls for customer requests and staff access. Search, analytics and AI-visibility execution remain separate professional services and are not bundled into this feature addendum.",
    )
    section(doc, "Privacy and staff access")
    b.add_table(
        doc,
        ["Area", "Added workflow", "Safeguard"],
        [
            ["Privacy requests", "Account-data export or deletion-review request", "Identity, retention and approval checks"],
            ["Sessions", "View/revoke sessions where enabled", "Device history and forced sign-out"],
            ["Suspicious access", "Security notification after defined events", "Rate limit, lockout and investigation trail"],
            ["Attachments", "Allowed support/return evidence", "Type/size validation, access and retention"],
            ["Exports", "Controlled business-data exports", "Role, reason and auditable download"],
        ],
        widths=[4.0, 12.0, 10.6],
        font_size=7.7,
    )
    section(doc, "Clients We Work With page")
    b.add_body(doc, "A new brand-trust page presents approved client or business names, logo marks, sectors or applications, and optional short testimonials or case studies. Client names, logos, quotations and performance claims require written permission and client-supplied approved assets. The existing content administration manages the page; no separate admin view is added.")
    section(doc, "Reserved for a separate proposal")
    b.add_table(
        doc,
        ["Search and measurement setup", "SEO and AI visibility service"],
        [[
            "Google Analytics 4 property and ecommerce configuration\nGoogle Tag Manager governance\nGoogle Search Console verification and monitoring\nGoogle Merchant Center onboarding\nMeasurement QA, dashboards and reporting",
            "SEO research and content strategy\nAEO and GEO strategy\nContent writing and publishing calendar\nSearch and AI-referral monitoring\nOngoing optimization and monthly reporting",
        ]],
        widths=[13.3, 13.3],
        font_size=8.25,
        header_fill=b.PURPLE,
    )
    b.add_body(doc, "Policies must be supplied or approved by the client and its advisers. MarketiX Studio implements the approved workflow but does not provide legal advice.", bold_lead="Client responsibility. ")

    b.new_page(
        doc,
        "07  Customer page inventory",
        "Fifteen additional customer destinations",
        "These destinations are added to the original sixty-seven customer-facing pages and screens, producing a revised customer count of eighty-two.",
    )
    customer_pages = [
        ["68", "Support Centre", "Guided self-service categories and support entry"],
        ["69", "My Tickets", "Ticket list, filters and status"],
        ["70", "Ticket Conversation", "Thread, attachments and linked records"],
        ["71", "Return and Exchange Eligibility", "Policy and item-level eligibility"],
        ["72", "Return or Refund Case Detail", "RMA, pickup, QC and resolution timeline"],
        ["73", "Exchange Selection", "Replacement variant and price difference"],
        ["74", "COD Refund Details", "Verified refund destination and status"],
        ["75", "Product Comparison", "Comparable product and variant attributes"],
        ["76", "GSM and Use Case Finder", "Guided product recommendation"],
        ["77", "Product Questions and Answers", "Product questions and approved answers"],
        ["78", "Back in Stock Subscriptions", "Variant subscriptions and alerts"],
        ["79", "Store Credit", "Balance and transaction history"],
        ["80", "Referral Centre", "Referral link, eligibility and rewards"],
        ["81", "Privacy Requests", "Account export or deletion review"],
        ["82", "Clients We Work With", "Approved client brands, sectors and references"],
    ]
    b.add_table(doc, ["New number", "Destination", "Purpose"], customer_pages, widths=[2.5, 8.4, 15.7], font_size=7.55)

    b.new_page(
        doc,
        "08  Administration inventory",
        "Service, after sales and delivery views",
        "The first fifteen additions establish operational ownership for support cases, returns, exchanges, refunds and delivery exceptions.",
    )
    admin_1 = [
        ["45", "Service Desk Dashboard", "Open volume, SLA, ageing and ownership"],
        ["46", "Ticket Queue", "Search, priority, filters and assignment"],
        ["47", "Ticket Detail", "Conversation, notes and linked records"],
        ["48", "SLA and Escalation Rules", "Response targets and overdue routing"],
        ["49", "Reply Templates", "Approved macros and version history"],
        ["50", "Support Analytics", "Response, resolution, reopen and satisfaction"],
        ["51", "Return Rules", "Eligibility and resolution policy"],
        ["52", "After Sales Queue", "Return, exchange, refund and replacement workload"],
        ["53", "Return Case and QC", "Evidence, pickup, inspection and disposition"],
        ["54", "Exchange Processing", "Replacement and price difference"],
        ["55", "Refund Approval", "Approval level and online refund control"],
        ["56", "COD Refund Workflow", "Destination, maker-checker and status"],
        ["57", "Store Credit Ledger", "Issue, redeem, expire and reverse"],
        ["58", "NDR Queue", "Failed delivery cases and deadlines"],
        ["59", "NDR Case", "Customer response and courier action history"],
    ]
    b.add_table(doc, ["New number", "Administration view", "Purpose"], admin_1, widths=[2.5, 8.5, 15.6], font_size=7.45)

    b.new_page(
        doc,
        "09  Administration inventory",
        "Trust, discovery, communication and governance views",
        "The remaining fourteen administration additions complete the revised inventory of seventy-three-plus major views. The additional Clients We Work With destination uses the existing content administration.",
    )
    admin_2 = [
        ["60", "RTO Risk Rules", "Pincode, order, customer and payment controls"],
        ["61", "RTO Analytics", "Cost, reason, carrier and recovery reporting"],
        ["62", "Review Moderation", "Verified status, media, decision and reply"],
        ["63", "Q&A Moderation", "Assignment, answer and publishing"],
        ["64", "Search Synonyms", "Approved query language and normalization"],
        ["65", "Search Merchandising", "Query, category and campaign rules"],
        ["66", "Search Analytics", "Queries, zero results, clicks and conversion"],
        ["67", "Comparison Configuration", "Comparable fields and products"],
        ["68", "GSM Finder Rules", "Questions, rules and recommendations"],
        ["69", "Notification Templates", "Versions, previews and approval"],
        ["70", "Notification Delivery Logs", "Provider outcome, retry and replay"],
        ["71", "Operations Exception Queue", "Payment, refund, shipment and job failures"],
        ["72", "Privacy Requests", "Verification, retention and resolution"],
        ["73", "Session and Device Security", "Sessions, revocation and risk events"],
    ]
    b.add_table(doc, ["New number", "Administration view", "Purpose"], admin_2, widths=[2.5, 8.5, 15.6], font_size=7.45)

    b.new_page(
        doc,
        "10  Delivery and acceptance",
        "Every expansion module is tested before dependent work proceeds",
        "Each selected module moves through requirements, UX approval, controlled build, integration testing, operational testing and client acceptance before production activation.",
    )
    b.add_table(
        doc,
        ["Gate", "Required output", "Acceptance evidence"],
        [
            ["Rules and states", "Workflow, roles, notifications, exceptions and policy dependencies", "Written criteria and state diagram"],
            ["UX approval", "Desktop/mobile, errors, empty/loading states and both themes", "Approved screens and interaction notes"],
            ["Controlled build", "Customer, admin, service and data workflow", "Developer checks and traceable issue list"],
            ["Integration test", "Sandbox payment, shipping and communication behaviour", "Success, failure, retry and duplicate-event tests"],
            ["Operational test", "Agent, warehouse, finance and manager scenarios", "Role-based test evidence and corrected defects"],
            ["Client acceptance", "Complete agreed module in staging", "Client sign-off before production"],
            ["Launch verification", "Production configuration and smoke test", "Health checks, sample workflow and rollback readiness"],
        ],
        widths=[4.0, 12.5, 10.1],
        font_size=7.85,
    )
    section(doc, "Minimum scenario coverage")
    b.add_body(doc, "Happy path, validation error, permission denial, provider timeout, duplicate webhook, customer retry, partial completion, cancellation, message failure, mobile layout, light mode, dark mode and audit-history verification.")
    section(doc, "Delivery boundary")
    b.add_body(doc, "Third-party onboarding, API availability and client policy approvals can affect sequencing. A module is accepted against its approved criteria rather than an assumed or undocumented workflow.")

    b.new_page(
        doc,
        "11  Deliverables and decisions",
        "The approved addendum becomes part of the controlled implementation plan",
        "Kleawip may approve all additions or divide them between launch and a later phase. The exact selected modules, dependencies and acceptance criteria will be recorded before development begins.",
    )
    section(doc, "Deliverables")
    b.add_bullets(doc, [
        "Approved workflow and acceptance criteria for every selected module.",
        "Responsive customer experiences in established light and dark modes.",
        "Role-based administration views with audit history for critical actions.",
        "Connected customer, order, payment, shipment, return and support records.",
        "Provider-aware failure states, retries and operations exception handling.",
        "Module-wise functional, responsive, permission and integration testing.",
        "Updated administration guide and client handover for added operations.",
        "Updated platform inventory of 82 customer screens and 73+ admin views.",
    ], columns=2, size=8.5)
    section(doc, "Client confirmations")
    b.add_table(
        doc,
        ["Decision", "Confirmation required"],
        [
            ["Module selection", "Approve all additions or identify launch and later phases"],
            ["Bulk purchase", "Quote-only or online bulk checkout, MOQ, quantity tiers, GST and fulfilment terms"],
            ["Client showcase", "Approved client names/logos, written permissions and approved references"],
            ["Support", "Roles, response targets, channels and escalation ownership"],
            ["After sales", "Return windows, exchange, COD refund and store-credit rules"],
            ["Delivery risk", "NDR actions, RTO thresholds and partial-COD restrictions"],
            ["Retention", "Back-in-stock, abandoned-cart, store-credit and referral choices"],
            ["Approvals", "Authorised staff for refunds, write-offs, exports and sensitive actions"],
        ],
        widths=[5.7, 20.9],
        font_size=7.85,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("MARKETIX STUDIO")
    b.set_repeatable_font(r, size=9.4, bold=True, color=RGBColor.from_string(b.PURPLE))
    r = p.add_run("   marketixstudio.com  |  contact@marketixstudio.com  |  +91 90217 53876")
    b.set_repeatable_font(r, size=8.8, color=b.MUTED)

    core = doc.core_properties
    core.title = "Kleawip Feature Expansion Addendum"
    core.subject = "Additional ecommerce capabilities beyond the original scope"
    core.author = "MarketiX Studio"
    core.keywords = "Kleawip, ecommerce, feature expansion, support, refunds, NDR, RTO"
    doc.save(b.OUT)
    print(b.OUT)


if __name__ == "__main__":
    build()
